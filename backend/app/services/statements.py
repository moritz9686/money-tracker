"""In-memory bank-statement parsers. Raw uploads are never persisted or logged."""

import csv
import io
import re
from abc import ABC, abstractmethod
from dataclasses import dataclass
from datetime import datetime, timezone
from decimal import Decimal, InvalidOperation
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from pypdf import PdfReader
from pypdf.errors import PyPdfError

from app.db.models import PaymentMode, TransactionType
from app.services.deduplication import normalize_merchant


class UnsupportedStatementError(ValueError):
    """Raised without including uploaded statement content."""


MAXIMUM_STATEMENT_BYTES = 15 * 1024 * 1024


@dataclass(frozen=True)
class StatementTransaction:
    transaction_date: datetime
    amount: Decimal
    transaction_type: TransactionType
    description: str
    reference_id: str | None
    account_last4: str | None
    merchant: str | None
    payment_mode: PaymentMode
    bank_name: str
    confidence: Decimal


class StatementParser(ABC):
    bank_name: str
    bank_markers: tuple[str, ...]

    def detects(self, text: str) -> bool:
        normalized = text.upper()
        return any(marker in normalized for marker in self.bank_markers)

    def parse_rows(self, rows: list[dict[str, str]]) -> list[StatementTransaction]:
        parsed: list[StatementTransaction] = []
        for row in rows:
            result = self._parse_row(
                {self._key(k): str(v).strip() for k, v in row.items()}
            )
            if result:
                parsed.append(result)
        return parsed

    @abstractmethod
    def _parse_row(self, row: dict[str, str]) -> StatementTransaction | None: ...

    def _common(self, row: dict[str, str]) -> StatementTransaction | None:
        date = self._date(self._first(row, "date", "transactiondate", "valuedate"))
        description = self._first(
            row, "description", "narration", "particulars", "remarks"
        )
        debit = self._money(self._first(row, "debit", "withdrawal", "debitamount"))
        credit = self._money(self._first(row, "credit", "deposit", "creditamount"))
        amount = debit or credit
        if date is None or amount is None or not description:
            return None
        direction = (
            TransactionType.DEBIT if debit is not None else TransactionType.CREDIT
        )
        reference = (
            self._first(row, "reference", "refno", "utr", "chqrefno", "transactionid")
            or None
        )
        account = re.search(
            r"(\d{4})$", self._first(row, "account", "accountnumber", "accountno")
        )
        text = description.upper()
        mode = (
            PaymentMode.UPI
            if "UPI" in text
            else PaymentMode.CARD
            if "CARD" in text
            else PaymentMode.NEFT
            if "NEFT" in text
            else PaymentMode.IMPS
            if "IMPS" in text
            else PaymentMode.OTHER
        )
        merchant = normalize_merchant(re.split(r"[/|]", description)[-1]) or None
        return StatementTransaction(
            date,
            amount,
            direction,
            description[:500],
            reference,
            account.group(1) if account else None,
            merchant,
            mode,
            self.bank_name,
            Decimal("0.900"),
        )

    @staticmethod
    def _key(value: str) -> str:
        return re.sub(r"[^a-z0-9]", "", value.lower())

    @staticmethod
    def _first(row: dict[str, str], *keys: str) -> str:
        return next((row.get(key, "") for key in keys if row.get(key, "")), "")

    @staticmethod
    def _money(value: str) -> Decimal | None:
        if not value or value.strip(" -") == "":
            return None
        try:
            return abs(Decimal(re.sub(r"[^0-9.-]", "", value).replace(",", "")))
        except InvalidOperation:
            return None

    @staticmethod
    def _date(value: str) -> datetime | None:
        for fmt in ("%d/%m/%Y", "%d-%m-%Y", "%Y-%m-%d", "%d/%m/%y"):
            try:
                return datetime.strptime(value.strip(), fmt).replace(
                    tzinfo=timezone.utc
                )
            except ValueError:
                pass
        return None


class HDFCParser(StatementParser):
    bank_name, bank_markers = "HDFC", ("HDFC BANK",)

    def _parse_row(self, row: dict[str, str]) -> StatementTransaction | None:
        return self._common(row)


class ICICIParser(StatementParser):
    bank_name, bank_markers = "ICICI", ("ICICI BANK",)

    def _parse_row(self, row: dict[str, str]) -> StatementTransaction | None:
        return self._common(row)


class SBIParser(StatementParser):
    bank_name, bank_markers = "SBI", ("STATE BANK", "SBI")

    def _parse_row(self, row: dict[str, str]) -> StatementTransaction | None:
        return self._common(row)


class AxisParser(StatementParser):
    bank_name, bank_markers = "Axis Bank", ("AXIS BANK",)

    def _parse_row(self, row: dict[str, str]) -> StatementTransaction | None:
        return self._common(row)


PARSERS: tuple[StatementParser, ...] = (
    HDFCParser(),
    ICICIParser(),
    SBIParser(),
    AxisParser(),
)


def parse_statement(
    content: bytes, filename: str
) -> tuple[StatementParser, list[StatementTransaction]]:
    if len(content) > MAXIMUM_STATEMENT_BYTES:
        raise UnsupportedStatementError("Statement exceeds the 15 MB size limit")
    suffix = filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    try:
        if suffix == "csv":
            text = content.decode("utf-8-sig", errors="replace")
            rows = list(csv.DictReader(io.StringIO(text)))
            sample = text[:2000]
        elif suffix == "xlsx":
            sheet = load_workbook(
                io.BytesIO(content), read_only=True, data_only=True
            ).active
            values = list(sheet.iter_rows(values_only=True))
            if not values:
                raise UnsupportedStatementError("Statement contains no table")
            headers = [str(v or "") for v in values[0]]
            rows = [
                dict(zip(headers, (str(v or "") for v in row), strict=False))
                for row in values[1:]
            ]
            sample = " ".join(headers)
        elif suffix == "pdf":
            text = "\n".join(
                page.extract_text() or ""
                for page in PdfReader(io.BytesIO(content)).pages
            )
            rows, sample = list(csv.DictReader(io.StringIO(text))), text[:2000]
        else:
            raise UnsupportedStatementError(
                "Only CSV, XLSX, and text-based PDF statements are supported"
            )
    except (BadZipFile, InvalidFileException, OSError, PyPdfError) as error:
        raise UnsupportedStatementError("Malformed or unreadable statement") from error
    parser = next((item for item in PARSERS if item.detects(sample)), None)
    if parser is None:
        raise UnsupportedStatementError("Unsupported bank statement")
    return parser, parser.parse_rows(rows)
